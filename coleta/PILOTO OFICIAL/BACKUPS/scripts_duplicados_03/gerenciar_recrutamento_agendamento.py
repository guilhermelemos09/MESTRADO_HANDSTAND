"""
SISTEMA OFICIAL DE GESTÃO DE RECRUTAMENTO, TRIAGEM, AGENDAMENTO E TCLE DIGITAL
Projeto de Mestrado: Análise de Fatores Associados ao Desempenho no Handstand e Handstand Walk
Mestrando: Guilherme de Paula Lemos | Orientador: Prof. Dr. Matheus Machado Gomes
Instituição: EEFERP - USP (2026)

Integração dos 2 Formulários Oficiais do Google:
1. Formulário 1 (Cadastro e Triagem): Contato, dados biológicos, disponibilidade e triagem clínica.
2. Formulário 2 (Dados da Pesquisa): Lateralidade (membro dominante), experiência, horas, modalidade e PR 1-RM Shoulder Press.
3. Consentimento Digital: Detecção de PDFs do TCLE e PAR-Q assinados na tela com caneta S-Pen em 'tcle_assinados/'.

Funcionalidades:
1. --status: Exibe resumo geral, participantes agendados, status do TCLE e próximo ID disponível (P002, P003...).
2. --listar: Lista candidatos no banco com filtros opcionais por dia da semana, turno e status.
3. --sincronizar: Lê novas respostas de ambos os formulários (via web sheets ou arquivos locais CSV/XLSX).
4. --agendar: Atribui o próximo ID sequencial oficial, agenda Sessão 1 e Sessão 2, atualiza a agenda visual e cria pastas de coleta.
5. --verificar-tcle: Varre a pasta 'tcle_assinados/' e dá baixa automática na Tabela Mestre para os arquivos encontrados.
6. --baixa-tcle "PID": Dá baixa manual imediata no TCLE do participante (ex: --baixa-tcle P002).
"""

import os
import sys
import csv
import glob
import re
import argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE_DIR = r"C:\Users\Gui\Documents\MESTRADO_HANDSTAND\coleta\PILOTO OFICIAL"
PASTA_00 = os.path.join(BASE_DIR, "00_CADASTRO_TRIAGEM_E_TCLE")
TABELA_MESTRE_PATH = os.path.join(PASTA_00, "TABELA_MESTRE_IDENTIFICACAO_PARTICIPANTES.xlsx")
if not os.path.exists(TABELA_MESTRE_PATH):
    TABELA_MESTRE_PATH = os.path.join(BASE_DIR, "TABELA_MESTRE_IDENTIFICACAO_PARTICIPANTES.xlsx")

PASTA_FORMS_1 = os.path.join(PASTA_00, "respostas_forms", "01_triagem_agendamento")
PASTA_FORMS_2 = os.path.join(PASTA_00, "respostas_forms", "02_perfil_pratica")
PASTA_TCLE_ASSINADOS = os.path.join(PASTA_00, "tcle_assinados")
CONFIG_URL_PATH = os.path.join(PASTA_00, "respostas_forms", "config_urls_forms.txt")

DADOS_BRUTOS_LACIDH = os.path.join(BASE_DIR, r"01_SESSAO_LACIDH_FORCA\dados_brutos")
PLANILHA_CAMPO_LACIDH = os.path.join(BASE_DIR, r"01_SESSAO_LACIDH_FORCA\PLANILHA_COLETA_CAMPO_LACIDH.xlsx")
PLANILHA_CAMPO_LABIOCOM = os.path.join(BASE_DIR, r"02_SESSAO_LABIOCOM_CINEMATICA\PLANILHA_COLETA_CAMPO_LABIOCOM.xlsx")

def normalizar_texto(txt):
    if txt is None:
        return ""
    return str(txt).strip()

def gerar_nome_abreviado(nome_completo):
    if not nome_completo:
        return ""
    partes = nome_completo.strip().split()
    if len(partes) > 1:
        return f"{partes[0]} {partes[-1]}"
    return nome_completo

def padronizar_telefone(tel_str):
    if not tel_str:
        return ""
    digitos = re.sub(r'\D', '', str(tel_str))
    if len(digitos) in [12, 13] and digitos.startswith("55"):
        digitos = digitos[2:]
    return digitos

def padronizar_idade_nascimento(val_str):
    if not val_str:
        return ""
    v_clean = str(val_str).strip()
    data_match = re.findall(r'\b(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\b', v_clean)
    if data_match:
        d, m, y = data_match[0]
        data_formatada = f"{int(d):02d}/{int(m):02d}/{y}"
        idade = 2026 - int(y)
        return f"{data_formatada} ({idade} anos)"
    return v_clean

def ler_config_urls():
    urls = {"form1": None, "form2": None}
    if os.path.exists(CONFIG_URL_PATH):
        try:
            with open(CONFIG_URL_PATH, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line.startswith("FORM1_SHEETS_EXPORT_CSV=") and len(line.split("=", 1)[1].strip()) > 5:
                        urls["form1"] = line.split("=", 1)[1].strip()
                    elif line.startswith("FORM2_SHEETS_EXPORT_CSV=") and len(line.split("=", 1)[1].strip()) > 5:
                        urls["form2"] = line.split("=", 1)[1].strip()
        except Exception as e:
            print(f"[AVISO] Erro ao ler config_urls_forms.txt: {e}")
    return urls

def baixar_respostas_nuvem(url_ou_id, destino_csv):
    import urllib.request
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url_ou_id)
    sheet_id = match.group(1) if match else url_ou_id.strip()
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
    
    try:
        req = urllib.request.Request(export_url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            content = resp.read()
            with open(destino_csv, "wb") as f:
                f.write(content)
        print(f"[OK] Planilha da nuvem baixada com sucesso: {os.path.basename(destino_csv)}")
        return True
    except Exception as e:
        print(f"[AVISO] Não foi possível baixar da nuvem ({e}). Verifique permissões públicas.")
        return False

def ler_arquivos_forms(pasta):
    arqs = glob.glob(os.path.join(pasta, "*.csv")) + glob.glob(os.path.join(pasta, "*.xlsx"))
    registros = []
    for arq in arqs:
        if "tabela_mestre" in os.path.basename(arq).lower():
            continue
        if arq.endswith(".csv"):
            try:
                with open(arq, "r", encoding="utf-8-sig") as f:
                    first_line = f.readline()
                    delim = ";" if ";" in first_line else ","
                with open(arq, "r", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f, delimiter=delim)
                    for row in reader:
                        registros.append(row)
            except Exception as e:
                print(f"[ERRO] Erro ao ler CSV {arq}: {e}")
        elif arq.endswith(".xlsx"):
            try:
                wb = openpyxl.load_workbook(arq, data_only=True)
                ws = wb.active
                headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
                for r in range(2, ws.max_row + 1):
                    row_dict = {}
                    for c, h in enumerate(headers, 1):
                        if h:
                            row_dict[str(h)] = str(ws.cell(r, c).value or "")
                    if any(row_dict.values()):
                        registros.append(row_dict)
                wb.close()
            except Exception as e:
                print(f"[ERRO] Erro ao ler XLSX {arq}: {e}")
    return registros

def mapear_dados_form1(row):
    res = {}
    for k, v in row.items():
        if not k:
            continue
        k_low = k.lower()
        v_clean = normalizar_texto(v)
        if "carimbo" in k_low or "timestamp" in k_low or "data/hora" in k_low:
            res["Carimbo_DataHora"] = v_clean
        elif "nome completo" in k_low or (k_low.startswith("nome") and "abreviado" not in k_low):
            if "Nome_Completo" not in res or len(v_clean) > len(res.get("Nome_Completo", "")):
                res["Nome_Completo"] = v_clean.rstrip(":").strip()
        elif "nascimento" in k_low or "data de nascimento" in k_low:
            res["Idade_Nascimento"] = padronizar_idade_nascimento(v_clean)
        elif "sexo" in k_low or "gênero" in k_low or "genero" in k_low:
            res["Sexo"] = "M" if "masc" in v_clean.lower() or v_clean.upper() == "M" else ("F" if "fem" in v_clean.lower() or v_clean.upper() == "F" else v_clean)
        elif "whatsapp" in k_low or ("telefone" in k_low and "emerg" not in k_low and "contato" in k_low) or ("telefone / whatsapp" in k_low):
            res["WhatsApp_Contato"] = padronizar_telefone(v_clean)
        elif "emergência" in k_low or "emergencia" in k_low:
            res["Telefone_Emergencia"] = padronizar_telefone(v_clean)
        elif "email" in k_low or "e-mail" in k_low:
            res["Email"] = v_clean
        elif "dor articular" in k_low or "lesão" in k_low or "lesao" in k_low or "limitação" in k_low:
            if "descreva" in k_low or "caso" in k_low:
                res["Descricao_Lesao"] = v_clean
            else:
                res["Historico_Lesao"] = "Sim" if "sim" in v_clean.lower() else "Não"
        elif "descreva brevemente" in k_low or "descreva sua condição" in k_low:
            res["Descricao_Lesao"] = v_clean
        elif "dia" in k_low or "dias" in k_low:
            res["Dias_Disponiveis"] = v_clean
        elif "turno" in k_low or "turnos" in k_low or "horário" in k_low or "horario" in k_low:
            res["Horarios_Preferencia"] = v_clean
        elif "sustentar" in k_low or "passos" in k_low or "critério" in k_low or "criterio" in k_low:
            res["Criterio_Triagem_HS"] = v_clean
        elif "modalidade" in k_low:
            res["Modalidade_Principal"] = v_clean
        elif "meses" in k_low and ("pratica" in k_low or "treina" in k_low):
            num = re.findall(r'\d+', v_clean)
            res["Tempo_Pratica_meses"] = num[0] if num else v_clean
        elif "volume" in k_low:
            num = re.findall(r'\d+', v_clean)
            res["Volume_Semanal_horas"] = num[0] if num else v_clean
    return res

def mapear_dados_form2(row):
    res = {}
    for k, v in row.items():
        if not k:
            continue
        k_low = k.lower()
        v_clean = normalizar_texto(v)
        if "carimbo" in k_low or "timestamp" in k_low:
            res["Carimbo_Perfil"] = v_clean
        elif "nome completo" in k_low or (k_low.startswith("nome") and "abreviado" not in k_low):
            res["Nome_Completo"] = v_clean.rstrip(":").strip()
        elif "dominante" in k_low or "mão dominante" in k_low or "membro superior dominante" in k_low:
            if "canhoto" in v_clean.lower() or "esquerdo" in v_clean.lower():
                res["Membro_Dominante"] = "E"
            elif "ambidestro" in v_clean.lower():
                res["Membro_Dominante"] = "A"
            else:
                res["Membro_Dominante"] = "D"
        elif "meses" in k_low and ("pratica" in k_low or "treina" in k_low):
            num = re.findall(r'\d+', v_clean)
            res["Tempo_Pratica_meses"] = num[0] if num else v_clean
        elif "volume" in k_low and ("semana" in k_low or "horas" in k_low or "médio" in k_low or "medio" in k_low):
            num = re.findall(r'\d+', v_clean)
            res["Volume_Semanal_horas"] = num[0] if num else v_clean
        elif "frequência" in k_low or "frequencia" in k_low:
            res["Frequencia_Semanal"] = v_clean
        elif "modalidade" in k_low and "principal" in k_low:
            res["Modalidade_Principal"] = v_clean
        elif "modalidade" in k_low and ("secundária" in k_low or "secundaria" in k_low or "outra" in k_low):
            res["Modalidades_Secundarias"] = v_clean
        elif "shoulder press" in k_low and ("carga" in k_low or "aproximada" in k_low or "kg" in k_low):
            num = re.findall(r'\d+(?:[.,]\d+)?', v_clean)
            res["PR_ShoulderPress_kg"] = num[0].replace(",", ".") if num else v_clean
        elif "apto" in k_low or "sem dor" in k_low or "dor articular" in k_low:
            res["Aptidao_Confirmada"] = "Sim" if "sim" in v_clean.lower() or "100%" in v_clean else "Não"
    return res

def sincronizar_formularios():
    print("=" * 80)
    print("SINCRONIZAÇÃO OFICIAL: GOOGLE FORMS (1 & 2) -> TABELA MESTRE (.XLSX)")
    print("=" * 80)

    urls = ler_config_urls()
    if urls.get("form1"):
        csv_dest = os.path.join(PASTA_FORMS_1, "respostas_nuvem_form1.csv")
        baixar_respostas_nuvem(urls["form1"], csv_dest)
    if urls.get("form2"):
        csv_dest = os.path.join(PASTA_FORMS_2, "respostas_nuvem_form2.csv")
        baixar_respostas_nuvem(urls["form2"], csv_dest)

    raw_f1 = ler_arquivos_forms(PASTA_FORMS_1)
    raw_f2 = ler_arquivos_forms(PASTA_FORMS_2)
    print(f"Respostas brutas carregadas -> Form 1 (Triagem): {len(raw_f1)} | Form 2 (Perfil): {len(raw_f2)}")

    f1_map = {}
    for r in raw_f1:
        d = mapear_dados_form1(r)
        nome = d.get("Nome_Completo", "")
        if nome:
            f1_map[nome.strip().lower()] = d

    f2_map = {}
    for r in raw_f2:
        d = mapear_dados_form2(r)
        nome = d.get("Nome_Completo", "")
        if nome:
            f2_map[nome.strip().lower()] = d

    if not os.path.exists(TABELA_MESTRE_PATH):
        print(f"[ERRO] Tabela Mestre não encontrada em {TABELA_MESTRE_PATH}")
        return 0

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH)
    ws_banco = wb["01_Banco_Interessados_Forms"]
    ws_oficiais = wb["02_Participantes_Oficiais"]

    existentes = {}
    for r in range(2, ws_banco.max_row + 1):
        n = normalizar_texto(ws_banco.cell(r, 2).value).lower()
        if n:
            existentes[n] = r

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    fill_novo = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")

    novos = 0
    atualizados = 0

    todos_nomes_forms = set(f1_map.keys()) | set(f2_map.keys())

    for n_lower in todos_nomes_forms:
        d1 = f1_map.get(n_lower, {})
        d2 = f2_map.get(n_lower, {})

        nome = d1.get("Nome_Completo") or d2.get("Nome_Completo", "")
        if not nome:
            continue

        mod = d2.get("Modalidade_Principal") or d1.get("Modalidade_Principal", "")
        tempo_hs = d2.get("Tempo_Pratica_meses") or d1.get("Tempo_Pratica_meses", "")
        vol_hs = d2.get("Volume_Semanal_horas") or d1.get("Volume_Semanal_horas", "")
        pr_sp = d2.get("PR_ShoulderPress_kg", "")
        membro_dom = d2.get("Membro_Dominante", "")
        criterio_hs = d1.get("Criterio_Triagem_HS", "")

        obs_itens = ["Triado via Formulários Oficiais."]
        if criterio_hs:
            obs_itens.append(f"Critério HS: {criterio_hs}")
        if membro_dom:
            obs_itens.append(f"Membro Dom: {membro_dom}")
        if pr_sp:
            obs_itens.append(f"PR-SP: {pr_sp}kg")
        if d1.get("Descricao_Lesao"):
            obs_itens.append(f"Obs Lesão: {d1['Descricao_Lesao']}")
        obs_texto = " | ".join(obs_itens)

        if n_lower in existentes:
            row_idx = existentes[n_lower]
            if mod and ws_banco.cell(row_idx, 9).value != mod:
                ws_banco.cell(row_idx, 9, value=mod)
                atualizados += 1
            if tempo_hs and ws_banco.cell(row_idx, 10).value != tempo_hs:
                ws_banco.cell(row_idx, 10, value=tempo_hs)
                atualizados += 1
            if vol_hs and ws_banco.cell(row_idx, 11).value != vol_hs:
                ws_banco.cell(row_idx, 11, value=vol_hs)
                atualizados += 1
            if d1.get("Dias_Disponiveis") and ws_banco.cell(row_idx, 12).value in ["A verificar", None, ""]:
                ws_banco.cell(row_idx, 12, value=d1["Dias_Disponiveis"])
                atualizados += 1
            if d1.get("Horarios_Preferencia") and ws_banco.cell(row_idx, 13).value in ["A verificar", None, ""]:
                ws_banco.cell(row_idx, 13, value=d1["Horarios_Preferencia"])
                atualizados += 1
            if obs_texto:
                ws_banco.cell(row_idx, 19, value=obs_texto)
            continue

        prox_linha = ws_banco.max_row + 1
        nova_linha = [
            d1.get("Carimbo_DataHora", d2.get("Carimbo_Perfil", "")),
            nome,
            gerar_nome_abreviado(nome),
            d1.get("Idade_Nascimento", ""),
            d1.get("Sexo", ""),
            d1.get("WhatsApp_Contato", ""),
            d1.get("Telefone_Emergencia", ""),
            d1.get("Email", ""),
            mod,
            tempo_hs,
            vol_hs,
            d1.get("Dias_Disponiveis", "A verificar"),
            d1.get("Horarios_Preferencia", "A verificar"),
            d1.get("Historico_Lesao", "Não"),
            "Disponível para Agendamento",
            "",
            "",
            "",
            obs_texto
        ]
        ws_banco.append(nova_linha)
        for c in range(1, len(nova_linha) + 1):
            cell = ws_banco.cell(prox_linha, c)
            cell.font = Font(name="Calibri", size=11)
            cell.border = thin_border
            if c in [1, 4, 5, 6, 7, 10, 11, 15, 16, 17, 18]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if c == 15:
                cell.fill = fill_novo

        existentes[n_lower] = prox_linha
        novos += 1
        print(f"  [+] Novo voluntário registrado no banco: {nome}")

    wb.save(TABELA_MESTRE_PATH)
    wb.close()
    print(f"\nSincronização finalizada: {novos} novo(s) cadastrado(s), {atualizados} campo(s) atualizado(s).")
    
    # Atualiza planilhas de campo se participantes já estiverem agendados
    atualizar_planilhas_campo_com_perfil(f2_map)

    # Verifica TCLE assinados
    verificar_tcle_digital()
    return novos

def atualizar_planilhas_campo_com_perfil(f2_map):
    """
    Percorre os participantes agendados na Tabela Mestre e preenche Membro_Dominante
    e PR_ShoulderPress nas planilhas de campo LaCiDH e LaBioCoM se ainda não preenchidos.
    """
    if not os.path.exists(TABELA_MESTRE_PATH):
        return

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH, data_only=True)
    ws_of = wb["02_Participantes_Oficiais"]
    
    wb_c = openpyxl.load_workbook(PLANILHA_CAMPO_LACIDH) if os.path.exists(PLANILHA_CAMPO_LACIDH) else None
    wb_lb = openpyxl.load_workbook(PLANILHA_CAMPO_LABIOCOM) if os.path.exists(PLANILHA_CAMPO_LABIOCOM) else None

    mod_c = False
    mod_lb = False

    for r in range(2, ws_of.max_row + 1):
        pid = ws_of.cell(r, 1).value
        nome = str(ws_of.cell(r, 2).value or "").strip()
        if not pid or not nome:
            continue
        
        d2 = f2_map.get(nome.lower(), {})
        if not d2:
            continue

        membro_dom = d2.get("Membro_Dominante", "")
        pr_sp = d2.get("PR_ShoulderPress_kg", "")

        m_num = re.findall(r'\d+', str(pid))
        if not m_num:
            continue
        linha_campo = int(m_num[0]) + 2

        if wb_c:
            ws_c = wb_c["Coleta_LaCiDH"]
            h_map_c = {str(c.value or "").strip(): idx for idx, c in enumerate(ws_c[1], start=1)}
            col_dom = h_map_c.get("Membro_Dominante", 15)
            col_pr = h_map_c.get("PR_Conhecida_1RM_kg", 25)
            if membro_dom and not ws_c.cell(linha_campo, col_dom).value:
                ws_c.cell(linha_campo, col_dom, value=membro_dom)
                mod_c = True
                print(f"  [CAMPO LaCiDH] Membro Dominante '{membro_dom}' atualizado para {pid} ({nome}).")
            if pr_sp and not ws_c.cell(linha_campo, col_pr).value:
                try:
                    ws_c.cell(linha_campo, col_pr, value=float(pr_sp))
                except:
                    ws_c.cell(linha_campo, col_pr, value=pr_sp)
                mod_c = True
                print(f"  [CAMPO LaCiDH] PR Shoulder Press '{pr_sp}kg' atualizado para {pid} ({nome}).")

        if wb_lb:
            ws_lb = wb_lb["Coleta_LaBioCoM"]
            if membro_dom and not ws_lb.cell(linha_campo, 5).value:
                ws_lb.cell(linha_campo, 5, value=membro_dom)
                mod_lb = True
                print(f"  [CAMPO LaBioCoM] Membro Dominante '{membro_dom}' atualizado para {pid} ({nome}).")

    wb.close()
    if wb_c:
        if mod_c:
            wb_c.save(PLANILHA_CAMPO_LACIDH)
        wb_c.close()
    if wb_lb:
        if mod_lb:
            wb_lb.save(PLANILHA_CAMPO_LABIOCOM)
        wb_lb.close()

def verificar_tcle_digital():
    """
    Varre a pasta tcle_assinados/ e atualiza o status na Tabela Mestre se houver PDFs assinados com caneta.
    """
    if not os.path.exists(PASTA_TCLE_ASSINADOS) or not os.path.exists(TABELA_MESTRE_PATH):
        return

    arquivos = glob.glob(os.path.join(PASTA_TCLE_ASSINADOS, "*.pdf"))
    if not arquivos:
        return

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH)
    ws_oficiais = wb["02_Participantes_Oficiais"]

    pids_assinados = set()
    for arq in arquivos:
        base_name = os.path.basename(arq).upper()
        m = re.search(r'(P\d{3})', base_name)
        if m:
            pids_assinados.add(m.group(1))

    atualizados = 0
    for r in range(2, ws_oficiais.max_row + 1):
        pid = ws_oficiais.cell(r, 1).value
        if pid in pids_assinados:
            status_atual = str(ws_oficiais.cell(r, 9).value or "")
            if "sim" not in status_atual.lower() and "concluída" not in status_atual.lower():
                ws_oficiais.cell(r, 9, value="Sim (Digital S-Pen)")
                atualizados += 1
                print(f"[TCLE DIGITAL] Baixa registrada para {pid} via arquivo assinado.")

    if atualizados > 0:
        wb.save(TABELA_MESTRE_PATH)
        print(f"[OK] {atualizados} participante(s) atualizados com TCLE Digital validado.")
    wb.close()

def dar_baixa_manual_tcle(pid_alvo):
    if not os.path.exists(TABELA_MESTRE_PATH):
        print("[ERRO] Tabela Mestre não encontrada.")
        return False

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH)
    ws_oficiais = wb["02_Participantes_Oficiais"]

    achou = False
    for r in range(2, ws_oficiais.max_row + 1):
        pid = str(ws_oficiais.cell(r, 1).value or "").strip().upper()
        if pid == pid_alvo.upper():
            nome = ws_oficiais.cell(r, 2).value
            ws_oficiais.cell(r, 9, value="Sim (Digital S-Pen)")
            achou = True
            print(f"[SUCESSO] Baixa manual de TCLE e PAR-Q confirmada para {pid} ({nome})!")
            break

    if achou:
        wb.save(TABELA_MESTRE_PATH)
    else:
        print(f"[AVISO] Participante {pid_alvo} não encontrado na lista oficial.")
    wb.close()
    return achou

def obter_status_recrutamento():
    if not os.path.exists(TABELA_MESTRE_PATH):
        print(f"Tabela Mestre não encontrada em: {TABELA_MESTRE_PATH}")
        return

    verificar_tcle_digital()

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH, data_only=True)
    ws_banco = wb["01_Banco_Interessados_Forms"]
    ws_oficiais = wb["02_Participantes_Oficiais"]

    total_banco = max(0, ws_banco.max_row - 1)
    agendados = 0
    coletados = 0
    proximo_id = "P001"

    oficiais_cadastrados = []
    for r in range(2, ws_oficiais.max_row + 1):
        pid = ws_oficiais.cell(r, 1).value
        nome = ws_oficiais.cell(r, 2).value
        s1 = ws_oficiais.cell(r, 7).value
        s2 = ws_oficiais.cell(r, 8).value
        status = str(ws_oficiais.cell(r, 9).value or "")

        if pid:
            if nome:
                oficiais_cadastrados.append((pid, nome, s1, s2, status))
                if "concluída" in status.lower() or "coletado" in status.lower():
                    coletados += 1
                else:
                    agendados += 1
            elif proximo_id == "P001":
                proximo_id = pid

    print("=" * 80)
    print("PAINEL GERAL DE RECRUTAMENTO, TRIAGEM E COLETA (PILOTO OFICIAL)")
    print("=" * 80)
    print(f"• Tabela Mestre Oficial:               {os.path.basename(TABELA_MESTRE_PATH)}")
    print(f"• Local:                               {os.path.dirname(TABELA_MESTRE_PATH)}")
    print(f"• Total de Voluntários no Banco:       {total_banco}")
    print(f"• Participantes com Coletas Concluídas:{coletados}")
    print(f"• Participantes Agendados / Em Curso:  {agendados}")
    print(f"• PRÓXIMO ID OFICIAL DISPONÍVEL:       >>> {proximo_id} <<<")
    print("-" * 80)
    print("PARTICIPANTES OFICIAIS REGISTRADOS:")
    if oficiais_cadastrados:
        for p in oficiais_cadastrados:
            print(f"  [{p[0]}] {p[1]} | S1: {p[2]} | S2: {p[3]} | Status: {p[4]}")
    else:
        print("  (Nenhum participante registrado ainda)")
    print("=" * 80)

def listar_interessados(filtro_dia=None, filtro_turno=None, status_filtro=None):
    if not os.path.exists(TABELA_MESTRE_PATH):
        print("Tabela Mestre não encontrada.")
        return

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH, data_only=True)
    ws_banco = wb["01_Banco_Interessados_Forms"]

    print("=" * 85)
    print("BANCO DE INTERESSADOS / VOLUNTÁRIOS DA TRIAGEM")
    print("=" * 85)

    encontrados = 0
    for r in range(2, ws_banco.max_row + 1):
        nome = ws_banco.cell(r, 2).value
        if not nome:
            continue
        tel = ws_banco.cell(r, 6).value or ""
        mod = ws_banco.cell(r, 9).value or ""
        dias = str(ws_banco.cell(r, 12).value or "")
        horarios = str(ws_banco.cell(r, 13).value or "")
        status = str(ws_banco.cell(r, 15).value or "")
        pid = ws_banco.cell(r, 16).value or ""

        if filtro_dia and filtro_dia.lower() not in dias.lower():
            continue
        if filtro_turno and filtro_turno.lower() not in horarios.lower():
            continue
        if status_filtro and status_filtro.lower() not in status.lower():
            continue

        encontrados += 1
        id_tag = f"[{pid}] " if pid else "[Disponível] "
        print(f"{encontrados}. {id_tag}{nome} | Tel: {tel} | Modalidade: {mod}")
        print(f"   Dias: {dias} | Horários: {horarios}")
        print(f"   Status: {status}")
        print("-" * 85)

    if encontrados == 0:
        print("Nenhum voluntário encontrado com os filtros selecionados.")
    print(f"Total listado: {encontrados}")

def agendar_participante(nome_busca, data_s1, hora_s1="08:00", data_s2=None, hora_s2="08:00"):
    if not data_s2:
        try:
            import datetime
            d1_obj = datetime.datetime.strptime(data_s1, "%d/%m/%Y")
            d2_obj = d1_obj + datetime.timedelta(days=7)
            data_s2 = d2_obj.strftime("%d/%m/%Y")
            print(f"[INFO] Data da Sessão 2 definida para janela ideal (7 dias após S1): {data_s2} às {hora_s2}")
        except Exception:
            data_s2 = "A definir"

    wb = openpyxl.load_workbook(TABELA_MESTRE_PATH)
    ws_banco = wb["01_Banco_Interessados_Forms"]
    ws_oficiais = wb["02_Participantes_Oficiais"]
    ws_agenda = wb["03_Agenda_Setembro_2026"]

    # 1. Localizar no banco
    linha_banco = None
    candidato = {}
    for r in range(2, ws_banco.max_row + 1):
        n = str(ws_banco.cell(r, 2).value or "").strip()
        if nome_busca.lower() in n.lower():
            linha_banco = r
            candidato = {
                "Nome_Completo": n,
                "Nome_Abreviado": ws_banco.cell(r, 3).value,
                "Idade_Nascimento": ws_banco.cell(r, 4).value,
                "Sexo": ws_banco.cell(r, 5).value,
                "Telefone": ws_banco.cell(r, 6).value,
                "Email": ws_banco.cell(r, 8).value,
                "Modalidade": ws_banco.cell(r, 9).value,
                "Tempo_Pratica": ws_banco.cell(r, 10).value,
                "Volume_Horas": ws_banco.cell(r, 11).value
            }
            break

    if not linha_banco:
        print(f"[ERRO] Candidato contendo '{nome_busca}' não encontrado no Banco de Interessados.")
        return False

    # 2. Localizar próximo ID vago
    proximo_id = None
    linha_oficial = None
    for r in range(2, ws_oficiais.max_row + 1):
        pid = ws_oficiais.cell(r, 1).value
        nome_atual = ws_oficiais.cell(r, 2).value
        if pid and not nome_atual:
            proximo_id = pid
            linha_oficial = r
            break

    if not proximo_id:
        print("[ERRO] Não há vagas ou IDs livres na tabela de participantes oficiais.")
        return False

    print(f"[+] Atribuindo ID Oficial: {proximo_id} para {candidato['Nome_Completo']}")

    # 3. Atualizar 02_Participantes_Oficiais
    ws_oficiais.cell(linha_oficial, 2, value=candidato["Nome_Completo"])
    ws_oficiais.cell(linha_oficial, 3, value=candidato["Nome_Abreviado"])
    ws_oficiais.cell(linha_oficial, 4, value=candidato["Idade_Nascimento"])
    ws_oficiais.cell(linha_oficial, 5, value=candidato["Sexo"])
    ws_oficiais.cell(linha_oficial, 6, value=candidato["Modalidade"])
    ws_oficiais.cell(linha_oficial, 7, value=data_s1)
    ws_oficiais.cell(linha_oficial, 8, value=data_s2)
    ws_oficiais.cell(linha_oficial, 9, value="Aguardando Assinatura")
    ws_oficiais.cell(linha_oficial, 10, value="Agendado - Aguardando Sessão 1")
    ws_oficiais.cell(linha_oficial, 11, value=candidato["Telefone"])
    ws_oficiais.cell(linha_oficial, 12, value=candidato["Email"])
    ws_oficiais.cell(linha_oficial, 13, value=f"01_SESSAO_LACIDH_FORCA\\dados_brutos\\{proximo_id}")

    # 4. Atualizar Banco
    ws_banco.cell(linha_banco, 15, value="Agendado Oficialmente")
    ws_banco.cell(linha_banco, 16, value=proximo_id)
    ws_banco.cell(linha_banco, 17, value=data_s1)
    ws_banco.cell(linha_banco, 18, value=data_s2)
    fill_agendado = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    ws_banco.cell(linha_banco, 15).fill = fill_agendado

    # 5. Inserir na Agenda Visual
    def marcar_agenda(data_str, hora_str, texto_slot, cor_bg):
        partes_data = data_str.split("/")
        if len(partes_data) == 3:
            col_match = f"{partes_data[0]}.{partes_data[1]}.{partes_data[2][-2:]}"
        else:
            col_match = data_str

        col_idx = None
        for c in range(2, ws_agenda.max_column + 1):
            val = str(ws_agenda.cell(1, c).value or "").strip()
            if col_match in val or data_str in val:
                col_idx = c
                break

        if col_idx:
            hora_num = int(hora_str.split(":")[0]) if ":" in hora_str else int(hora_str)
            linha_horario = None
            for r in range(3, ws_agenda.max_row + 1):
                val_h = ws_agenda.cell(r, 1).value
                if isinstance(val_h, (int, float)):
                    h_calc = round(val_h * 24)
                    if h_calc == hora_num:
                        linha_horario = r
                        break
                elif str(hora_num) in str(val_h):
                    linha_horario = r
                    break

            if not linha_horario:
                linha_horario = 4

            cell = ws_agenda.cell(linha_horario, col_idx)
            cell.value = texto_slot
            cell.fill = PatternFill(start_color=cor_bg, end_color=cor_bg, fill_type="solid")
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    marcar_agenda(data_s1, hora_s1, f"{proximo_id} - S1", "FEF08A")
    marcar_agenda(data_s2, hora_s2, f"{proximo_id} - S2", "BAE6FD")

    wb.save(TABELA_MESTRE_PATH)
    wb.close()

    # 6. Criar pastas físicas
    pasta_part = os.path.join(DADOS_BRUTOS_LACIDH, proximo_id)
    os.makedirs(pasta_part, exist_ok=True)
    os.makedirs(os.path.join(pasta_part, "Dados_biodex"), exist_ok=True)

    # 7. Atualizar Planilhas de Campo
    id_num = int(proximo_id.replace("P", ""))
    linha_campo = id_num + 2

    # Recupera dados do Form 2 se disponíveis
    raw_f2 = ler_arquivos_forms(PASTA_FORMS_2)
    f2_map = {mapear_dados_form2(r).get("Nome_Completo", "").strip().lower(): mapear_dados_form2(r) for r in raw_f2 if mapear_dados_form2(r).get("Nome_Completo")}
    d2 = f2_map.get(candidato["Nome_Completo"].strip().lower(), {})
    membro_dom = d2.get("Membro_Dominante", "")
    pr_sp = d2.get("PR_ShoulderPress_kg", "")

    if os.path.exists(PLANILHA_CAMPO_LACIDH):
        try:
            wb_c = openpyxl.load_workbook(PLANILHA_CAMPO_LACIDH)
            ws_c = wb_c["Coleta_LaCiDH"]
            h_map_c = {str(c.value or "").strip(): idx for idx, c in enumerate(ws_c[1], start=1)}
            col_dom = h_map_c.get("Membro_Dominante", 15)
            col_pr = h_map_c.get("PR_Conhecida_1RM_kg", 25)
            ws_c.cell(linha_campo, 1, value=proximo_id)
            ws_c.cell(linha_campo, 2, value=candidato["Nome_Abreviado"])
            ws_c.cell(linha_campo, 3, value=data_s1)
            ws_c.cell(linha_campo, 4, value=candidato["Sexo"])
            m_id = re.search(r'\((\d+)\s*anos\)', str(candidato["Idade_Nascimento"]))
            if m_id:
                ws_c.cell(linha_campo, 5, value=int(m_id.group(1)))
            if membro_dom:
                ws_c.cell(linha_campo, col_dom, value=membro_dom)
            if pr_sp:
                try:
                    ws_c.cell(linha_campo, col_pr, value=float(pr_sp))
                except:
                    ws_c.cell(linha_campo, col_pr, value=pr_sp)
            wb_c.save(PLANILHA_CAMPO_LACIDH)
            wb_c.close()
        except Exception as e:
            print(f"[AVISO] Erro ao atualizar campo LaCiDH: {e}")

    if os.path.exists(PLANILHA_CAMPO_LABIOCOM):
        try:
            wb_lb = openpyxl.load_workbook(PLANILHA_CAMPO_LABIOCOM)
            ws_lb = wb_lb["Coleta_LaBioCoM"]
            ws_lb.cell(linha_campo, 1, value=proximo_id)
            ws_lb.cell(linha_campo, 2, value=candidato["Nome_Abreviado"])
            ws_lb.cell(linha_campo, 3, value=data_s2)
            if membro_dom:
                ws_lb.cell(linha_campo, 5, value=membro_dom)
            wb_lb.save(PLANILHA_CAMPO_LABIOCOM)
            wb_lb.close()
        except Exception as e:
            print(f"[AVISO] Erro ao atualizar campo LaBioCoM: {e}")

    print("=" * 80)
    print(f"[SUCESSO TOTAL] Participante agendado oficialmente como {proximo_id}!")
    print(f"• Nome: {candidato['Nome_Completo']} ({candidato['Nome_Abreviado']})")
    print(f"• Sessão 1 (LaCiDH):    {data_s1} às {hora_s1}")
    print(f"• Sessão 2 (LaBioCoM):  {data_s2} às {hora_s2}")
    print(f"• Pasta de Coleta:      {pasta_part}")
    print("=" * 80)
    return True

def main():
    parser = argparse.ArgumentParser(description="Gerenciador de Recrutamento, Triagem e TCLE Handstand")
    parser.add_argument("--status", action="store_true", help="Exibe resumo do recrutamento e próximo ID disponível")
    parser.add_argument("--listar", action="store_true", help="Lista interessados cadastrados no banco")
    parser.add_argument("--dia", type=str, help="Filtra por dia da semana (ex: quinta)")
    parser.add_argument("--turno", type=str, help="Filtra por turno (ex: manhã)")
    parser.add_argument("--sincronizar", action="store_true", help="Sincroniza respostas locais ou da nuvem com a Tabela Mestre")
    parser.add_argument("--agendar", type=str, help="Nome do candidato a ser agendado")
    parser.add_argument("--data1", type=str, help="Data da Sessão 1 (DD/MM/AAAA)")
    parser.add_argument("--hora1", type=str, default="09:00", help="Horário da Sessão 1 (ex: 09:00)")
    parser.add_argument("--data2", type=str, help="Data da Sessão 2 (DD/MM/AAAA)")
    parser.add_argument("--hora2", type=str, default="09:00", help="Horário da Sessão 2 (ex: 09:00)")
    parser.add_argument("--verificar-tcle", action="store_true", help="Varre a pasta tcle_assinados/ e dá baixa na Tabela Mestre")
    parser.add_argument("--baixa-tcle", type=str, help="Dá baixa manual no TCLE de um ID (ex: --baixa-tcle P002)")

    args = parser.parse_args()

    if args.status:
        obter_status_recrutamento()
    elif args.listar or args.dia or args.turno:
        listar_interessados(filtro_dia=args.dia, filtro_turno=args.turno)
    elif args.sincronizar:
        sincronizar_formularios()
    elif args.verificar_tcle:
        verificar_tcle_digital()
    elif args.baixa_tcle:
        dar_baixa_manual_tcle(args.baixa_tcle)
    elif args.agendar:
        if not args.data1:
            print("[ERRO] Para agendar, informe ao menos --data1 DD/MM/AAAA.")
            sys.exit(1)
        agendar_participante(args.agendar, args.data1, args.hora1, args.data2, args.hora2)
    else:
        obter_status_recrutamento()

if __name__ == "__main__":
    main()
