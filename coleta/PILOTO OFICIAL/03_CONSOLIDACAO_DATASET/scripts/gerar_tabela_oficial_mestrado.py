"""
GERADOR DA TABELA OFICIAL DE VARIÁVEIS E DATASET MASTER (SPSS / ESTATÍSTICA FINAL)
Projeto de Mestrado: Análise de Fatores Associados ao Desempenho no Handstand e Handstand Walk
Mestrando: Guilherme de Paula Lemos | Orientador: Prof. Dr. Matheus Machado Gomes
Baseado rigorosamente no Projeto de Comitê de Ética (31/08/2026)

Mapeia:
- Todos os 6 Modelos Temáticos de Regressão Linear e Logística do Handstand Estático (HS)
- Todos os 6 Modelos Temáticos de Regressão Linear do Handstand Walk Dinâmico (HSW)
- Todas as variáveis de Capacidade Física (LaCiDH), Bioimpedância (Sanny), Cinemática 3D (Vicon) e Cinética (Bertec)
- Pré-consolidação com os dados reais do primeiro participante coletado (P001)

Gera:
- TABELA_OFICIAL_MESTRADO_ESTATISTICA_FINAL.xlsx (Com 2 abas: Dataset SPSS e Dicionário de Modelos)
- TABELA_OFICIAL_MESTRADO_ESTATISTICA_FINAL.csv (Pronto para importação no SPSS)
"""

import os
import sys
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# Definição do Dicionário Completo de Variáveis (Ordem Metodológica)
VARIAVEIS_METRICAS = [
    # --- BLOCO 1: IDENTIFICAÇÃO ---
    {
        "nome": "ID_Participante",
        "sub": "ex: P001",
        "label": "Identificador Único do Participante",
        "unidade": "Texto",
        "bloco": "1. Identificação",
        "cor_hex": "203764", # Azul Escuro
        "modelos": "Controle de Amostra",
        "papel": "Identificador",
        "tipo_spss": "Nominal",
        "p001_val": "P001"
    },
    {
        "nome": "Nome_Completo",
        "sub": "Nome Completo",
        "label": "Nome Completo do Voluntário",
        "unidade": "Texto",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "Identificação Interna",
        "papel": "Metadado",
        "tipo_spss": "Nominal",
        "p001_val": "Gustavo Henrique Donato da Costa"
    },
    {
        "nome": "Nome_Abreviado",
        "sub": "Nome Abreviado",
        "label": "Nome Abreviado para Relatórios",
        "unidade": "Texto",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "Identificação Interna",
        "papel": "Metadado",
        "tipo_spss": "Nominal",
        "p001_val": "Gustavo Donato"
    },
    {
        "nome": "Data_Coleta_Sessao1",
        "sub": "DD/MM/AAAA",
        "label": "Data da Sessão 1 - Força e BIA (LaCiDH)",
        "unidade": "Data",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "Metadado Operacional",
        "papel": "Metadado",
        "tipo_spss": "Data",
        "p001_val": "02/09/2026"
    },
    {
        "nome": "Data_Coleta_Sessao2",
        "sub": "DD/MM/AAAA",
        "label": "Data da Sessão 2 - Cinemática e Cinética (LaBioCoM)",
        "unidade": "Data",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "Metadado Operacional",
        "papel": "Metadado",
        "tipo_spss": "Data",
        "p001_val": "09/09/2026"
    },
    {
        "nome": "Sexo",
        "sub": "M / F",
        "label": "Sexo Biológico",
        "unidade": "M ou F",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "Estatística Descritiva / Covariável",
        "papel": "Covariável / Descritiva",
        "tipo_spss": "Nominal",
        "p001_val": "M"
    },
    {
        "nome": "Idade_anos",
        "sub": "Anos",
        "label": "Idade Cronológica",
        "unidade": "Anos",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "HS: Modelo 5 (Perfil) | HSW: Modelo 6 (Perfil)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 20
    },
    {
        "nome": "Membro_Dominante",
        "sub": "D ou E",
        "label": "Membro Superior Dominante",
        "unidade": "D ou E",
        "bloco": "1. Identificação",
        "cor_hex": "203764",
        "modelos": "Referência de Assimetria e FFP",
        "papel": "Controle Metodológico",
        "tipo_spss": "Nominal",
        "p001_val": "D"
    },

    # --- BLOCO 2: PERFIL DE PRÁTICA ---
    {
        "nome": "Modalidade_Principal",
        "sub": "CrossFit / Ginástica / etc",
        "label": "Modalidade Esportiva Principal",
        "unidade": "Texto",
        "bloco": "2. Perfil e Prática",
        "cor_hex": "1F4E78", # Azul Petróleo
        "modelos": "Estatística Descritiva / Comparação entre Grupos",
        "papel": "Fator Descritivo",
        "tipo_spss": "Nominal",
        "p001_val": "Calistenia, Handbalancing"
    },
    {
        "nome": "Perfil_Pratica",
        "sub": "Exclusiva vs Multimodal",
        "label": "Perfil de Prática Esportiva",
        "unidade": "Categórica",
        "bloco": "2. Perfil e Prática",
        "cor_hex": "1F4E78",
        "modelos": "Estatística Descritiva",
        "papel": "Fator Descritivo",
        "tipo_spss": "Nominal",
        "p001_val": "Exclusiva"
    },
    {
        "nome": "Tempo_Pratica_Total_meses",
        "sub": "Meses",
        "label": "Tempo de Prática na Modalidade de Origem",
        "unidade": "Meses",
        "bloco": "2. Perfil e Prática",
        "cor_hex": "1F4E78",
        "modelos": "Estatística Descritiva",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 24
    },
    {
        "nome": "Tempo_Treino_Especifico_HS_meses",
        "sub": "Meses",
        "label": "Tempo de Treinamento Específico no Handstand",
        "unidade": "Meses",
        "bloco": "2. Perfil e Prática",
        "cor_hex": "1F4E78",
        "modelos": "HS: Modelo 6 (Regressão Logística - Autonomia)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 18
    },
    {
        "nome": "Volume_Pratica_Acumulada_horas",
        "sub": "Horas Totais",
        "label": "Volume de Prática Acumulada Estimada",
        "unidade": "Horas",
        "bloco": "2. Perfil e Prática",
        "cor_hex": "1F4E78",
        "modelos": "HS: Modelo 5 (Perfil) | HSW: Modelo 6 (Perfil)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 150.0
    },
    {
        "nome": "Frequencia_Semanal_treinos",
        "sub": "Dias/semana",
        "label": "Frequência Semanal de Treinamento",
        "unidade": "Dias/semana",
        "bloco": "2. Perfil e Prática",
        "cor_hex": "1F4E78",
        "modelos": "Estatística Descritiva",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 4
    },

    # --- BLOCO 3: ANTROPOMETRIA E BIOIMPEDÂNCIA ---
    {
        "nome": "Massa_Corporal_kg",
        "sub": "Balança LaCiDH (kg)",
        "label": "Massa Corporal Total (Balança)",
        "unidade": "kg",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723", # Verde Floresta
        "modelos": "Normalização de Força (1-RM, FPM, FFP)",
        "papel": "Variável de Ajuste Mecânico",
        "tipo_spss": "Escala",
        "p001_val": 65.25
    },
    {
        "nome": "Estatura_cm",
        "sub": "Estadiômetro (cm)",
        "label": "Estatura Total em Ortostase",
        "unidade": "cm",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Cálculo do IMC e Normalização Antropométrica",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 171.0
    },
    {
        "nome": "IMC_kg_m2",
        "sub": "kg/m²",
        "label": "Índice de Massa Corporal",
        "unidade": "kg/m²",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Caracterização Amostral",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 22.31
    },
    {
        "nome": "IMLG_FFMI_kg_m2",
        "sub": "FFMI (kg/m²)",
        "label": "Índice de Massa Livre de Gordura (Fat-Free Mass Index)",
        "unidade": "kg/m²",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Caracterização Muscular Avançada (Substituto do IMC para Atletas)",
        "papel": "Descritiva / Preditor de Robustez",
        "tipo_spss": "Escala",
        "p001_val": 18.60
    },
    {
        "nome": "Resistencia_BIA_ohms",
        "sub": "Sanny BIA1011AF (Ω)",
        "label": "Resistência Bioelétrica a 50 kHz",
        "unidade": "Ω",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Parâmetro Elétrico Bruto BIA",
        "papel": "Metadado Bioelétrico",
        "tipo_spss": "Escala",
        "p001_val": 497.9
    },
    {
        "nome": "Reatancia_BIA_ohms",
        "sub": "Sanny BIA1011AF (Ω)",
        "label": "Reatância Capacitiva a 50 kHz",
        "unidade": "Ω",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Parâmetro Elétrico Bruto BIA",
        "papel": "Metadado Bioelétrico",
        "tipo_spss": "Escala",
        "p001_val": 78.7
    },
    {
        "nome": "Impedancia_BIA_ohms",
        "sub": "Sanny BIA1011AF (Ω)",
        "label": "Impedância Vetorial (Z)",
        "unidade": "Ω",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Parâmetro Elétrico Bruto BIA",
        "papel": "Metadado Bioelétrico",
        "tipo_spss": "Escala",
        "p001_val": 504.1
    },
    {
        "nome": "Angulo_Fase_deg",
        "sub": "Graus (°)",
        "label": "Ângulo de Fase (Sanny)",
        "unidade": "Graus (°)",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Integridade Celular / Vitalidade Muscular",
        "papel": "Descritiva / Biomarcador",
        "tipo_spss": "Escala",
        "p001_val": 8.98
    },
    {
        "nome": "BIA_Nivel_Atividade",
        "sub": "Sanny (Classificação)",
        "label": "Nível de Atividade Física (BIA Sanny)",
        "unidade": "Categórica",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Caracterização Amostral (ACSM / OMS)",
        "papel": "Fator Descritivo",
        "tipo_spss": "Nominal",
        "p001_val": "Pouco Ativo"
    },
    {
        "nome": "Agua_Corporal_Total_kg",
        "sub": "ACT (kg)",
        "label": "Água Corporal Total",
        "unidade": "kg",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Composição Corporal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 39.05
    },
    {
        "nome": "Agua_Corporal_pct",
        "sub": "ACT (%)",
        "label": "Percentual de Água Corporal Total",
        "unidade": "%",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Composição Corporal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 59.89
    },
    {
        "nome": "Percentual_Gordura_pct",
        "sub": "% Gordura",
        "label": "Percentual de Gordura Corporal",
        "unidade": "%",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "HS: Modelo 4 (Antropometria) | HSW: Modelo 5 (Antropometria)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 16.56
    },
    {
        "nome": "Massa_Gorda_kg",
        "sub": "kg",
        "label": "Massa Gorda Total",
        "unidade": "kg",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Composição Corporal (Massa Inerte)",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 10.80
    },
    {
        "nome": "Massa_Livre_Gordura_kg",
        "sub": "Massa Magra (kg)",
        "label": "Massa Livre de Gordura / Massa Magra Total",
        "unidade": "kg",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "HS: Modelo 4 (Antropometria) | HSW: Modelo 5 (Antropometria)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 54.40
    },
    {
        "nome": "Massa_Muscular_Esqueletica_kg",
        "sub": "SMM (kg)",
        "label": "Massa Muscular Esquelética",
        "unidade": "kg",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Composição Corporal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 31.06
    },
    {
        "nome": "Taxa_Metabolica_Basal_kcal",
        "sub": "kcal",
        "label": "Taxa Metabólica Basal Estimada",
        "unidade": "kcal",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "Composição Corporal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 1656.5
    },
    {
        "nome": "HS_Altura_CoM_mm",
        "sub": "Vicon (mm)",
        "label": "Altura Vertical do Centro de Massa em Inversão",
        "unidade": "mm",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "HS: Modelo 4 (Antropometria) | HSW: Modelo 5 (Antropometria)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 1128.31
    },
    {
        "nome": "HS_Comprimento_Bracos_mm",
        "sub": "Vicon (mm)",
        "label": "Comprimento dos Membros Superiores (Alavanca)",
        "unidade": "mm",
        "bloco": "3. Antropometria & BIA",
        "cor_hex": "385723",
        "modelos": "HS: Modelo 4 (Antropometria) | HSW: Modelo 5 (Antropometria)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 511.05
    },

    # --- BLOCO 4: CAPACIDADE FÍSICA - SESSÃO 1 (LaCiDH) ---
    # FPM
    {
        "nome": "FPM_Dom_Max_kgf",
        "sub": "kgf",
        "label": "FPM Máxima no Membro Dominante",
        "unidade": "kgf",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11", # Laranja Queimado
        "modelos": "Força Isométrica Distal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 46.0
    },
    {
        "nome": "FPM_NaoDom_Max_kgf",
        "sub": "kgf",
        "label": "FPM Máxima no Membro Não-Dominante",
        "unidade": "kgf",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Força Isométrica Distal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 52.0
    },
    {
        "nome": "FPM_Max_Global_kgf",
        "sub": "kgf",
        "label": "FPM Máxima Global (Maior Pico entre Membros)",
        "unidade": "kgf",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Força Isométrica Distal Absoluta",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 52.0
    },
    {
        "nome": "FPM_Media_Bilateral_kgf",
        "sub": "kgf",
        "label": "FPM Média Bilateral dos Membros",
        "unidade": "kgf",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Força Isométrica Distal",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 49.0
    },
    {
        "nome": "FPM_Relativa_Media_kgf_kg",
        "sub": "kgf/kg",
        "label": "Força de Preensão Manual Relativa (Média Bilateral / MC)",
        "unidade": "kgf/kg",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "HS: Modelo 1 (Capacidade Física)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 0.7510
    },
    {
        "nome": "FPM_Relativa_Max_kgf_kg",
        "sub": "kgf/kg",
        "label": "Força de Preensão Manual Relativa Máxima (Pico Global / MC)",
        "unidade": "kgf/kg",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Alternativa de Força Distal Normalizada",
        "papel": "Preditor / Sensibilidade",
        "tipo_spss": "Escala",
        "p001_val": 0.7969
    },

    # Biodex FFP
    {
        "nome": "FFP_Biodex_Dom_Max_Nm",
        "sub": "N.m",
        "label": "Pico de Torque Isométrico Flexores Punho Dominante (70°)",
        "unidade": "N.m",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Dinamometria Isométrica Biodex PRO",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 18.5
    },
    {
        "nome": "FFP_Biodex_NaoDom_Max_Nm",
        "sub": "N.m",
        "label": "Pico de Torque Isométrico Flexores Punho Não-Dominante (70°)",
        "unidade": "N.m",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Dinamometria Isométrica Biodex PRO",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 20.5
    },
    {
        "nome": "FFP_Biodex_Media_Bilateral_Nm",
        "sub": "N.m",
        "label": "Média Bilateral do Pico de Torque dos Flexores de Punho",
        "unidade": "N.m",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Força Específica de Suporte",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 19.5
    },
    {
        "nome": "FFP_Biodex_Relativo_Media_Nm_kg",
        "sub": "N.m/kg",
        "label": "Pico de Torque Relativo Flexores Punho (Média Bilateral / MC)",
        "unidade": "N.m/kg",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "HS: Modelo 1 (Capacidade Física)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 0.2989
    },
    {
        "nome": "LSI_FFP_pct",
        "sub": "%",
        "label": "Índice de Simetria de Membros nos Flexores de Punho",
        "unidade": "%",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "HSW: Modelo 1 (Capacidade Física)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 110.81
    },

    # 1-RM Shoulder Press
    {
        "nome": "SP_1RM_Estagio_Final_kg",
        "sub": "kg",
        "label": "Carga Máxima de 1-RM no Shoulder Press (Estágio Final)",
        "unidade": "kg",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Força Máxima Dinâmica Absoluta",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 56.0
    },
    {
        "nome": "1RM_ShoulderPress_Relativa_kg_kg",
        "sub": "kg/kg",
        "label": "Força Relativa no Shoulder Press (1-RM / Massa Corporal)",
        "unidade": "kg/kg",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "HS: Modelo 1 (Física), Modelo 6 (Logístico) | HSW: Modelo 1 (Física)",
        "papel": "Preditor Independente Central",
        "tipo_spss": "Escala",
        "p001_val": 0.8582
    },

    # Wall-HS Resistência
    {
        "nome": "Resistencia_WallHS_T1_s",
        "sub": "Segundos (s)",
        "label": "Resistência Isométrica Wall-Supported HS Tentativa 1",
        "unidade": "s",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Tolerância à Fadiga em Inversão",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 60.0
    },
    {
        "nome": "Resistencia_WallHS_T2_s",
        "sub": "Segundos (s)",
        "label": "Resistência Isométrica Wall-Supported HS Tentativa 2",
        "unidade": "s",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "Tolerância à Fadiga em Inversão",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 64.0
    },
    {
        "nome": "Resistencia_WallHS_Max_s",
        "sub": "Segundos (s)",
        "label": "Tempo Máximo de Resistência Wall-Supported HS",
        "unidade": "s",
        "bloco": "4. Capacidade Física (LaCiDH)",
        "cor_hex": "C55A11",
        "modelos": "HS: Modelo 1 (Capacidade Física) | HSW: Modelo 1 (Capacidade Física)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 64.0
    },

    # --- BLOCO 5: BIOMECÂNICA HANDSTAND ESTÁTICO (HS - LaBioCoM) ---
    {
        "nome": "HS_Assist_Tempo_Sustentacao_Pico_s",
        "sub": "Segundos (s)",
        "label": "Melhor Tempo de Sustentação no HS Assistido (Pico)",
        "unidade": "s",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0", # Roxo Nobre
        "modelos": "DESFECHO PRINCIPAL HS (Modelos 1 a 5) | Preditor HS Mod 6 e HSW Mod 4",
        "papel": "Variável Dependente (Desfecho) / Preditor",
        "tipo_spss": "Escala",
        "p001_val": 40.80
    },
    {
        "nome": "HS_Livre_Sucesso_Binario",
        "sub": "1=Sim / 0=Não",
        "label": "Autonomia de Estabilização no HS Livre (≥ 3 segundos)",
        "unidade": "Dicotômica",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "DESFECHO MODELO 6 (Regressão Logística Binária HS)",
        "papel": "Variável Dependente (Desfecho Logístico)",
        "tipo_spss": "Nominal",
        "p001_val": 1
    },
    {
        "nome": "HS_Livre_Tempo_Max_s",
        "sub": "Segundos (s)",
        "label": "Tempo Máximo de Sustentação no HS Livre",
        "unidade": "s",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva HS",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 24.50
    },
    {
        "nome": "HS_Assist_Distancia_CoP_CoM_Media_mm",
        "sub": "Milímetros (mm)",
        "label": "Distância Média Contínua entre CoP e Projeção do CoM",
        "unidade": "mm",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 2 (Eficiência do Controle Postural)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 50.02
    },
    {
        "nome": "HS_Assist_Velocidade_CoP_mm_s",
        "sub": "mm/s",
        "label": "Velocidade Média de Deslocamento do Centro de Pressão",
        "unidade": "mm/s",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 2 (Eficiência do Controle Postural)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 98.98
    },
    {
        "nome": "HS_Assist_ApEn_CoP",
        "sub": "Unidades Arbitr.",
        "label": "Entropia Aproximada do Deslocamento do CoP",
        "unidade": "Adimensional",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 2 (Controle Postural) | HSW: Modelo 4 (Transferência)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 0.7756
    },
    {
        "nome": "HS_Assist_Indice_Verticalidade_deg",
        "sub": "Graus (°)",
        "label": "Índice de Verticalidade 3D (Alinhamento Ombro-Quadril)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 12.02
    },
    {
        "nome": "HS_Assist_Extensao_Cervical_deg",
        "sub": "Graus (°)",
        "label": "Ângulo de Extensão Cervical em Inversão",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 23.35
    },
    {
        "nome": "HS_Assist_Flexao_Plantar_deg",
        "sub": "Graus (°)",
        "label": "Ângulo de Flexão Plantar do Tornozelo em Inversão",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 33.45
    },
    {
        "nome": "HS_Assist_Base_Apoio_mm",
        "sub": "Milímetros (mm)",
        "label": "Largura Linear da Base de Apoio (Distância entre Punhos)",
        "unidade": "mm",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 570.41
    },
    {
        "nome": "HS_Assist_Rotacao_Maos_deg",
        "sub": "Graus (°)",
        "label": "Orientação Angular das Mãos (Rotação Externa)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "HS: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 30.93
    },

    # --- Estratégias Articulares de Busca do Equilíbrio (HS) ---
    {
        "nome": "HS_Cotovelo_ROM_deg",
        "sub": "Graus (°)",
        "label": "Amplitude Total do Ângulo do Cotovelo (ROM)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Estratégia Motora",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 81.65
    },
    {
        "nome": "HS_Cotovelo_SD_deg",
        "sub": "Graus (°)",
        "label": "Variabilidade Postural do Cotovelo (Desvio Padrão)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Estratégia de Cotovelo",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 16.96
    },
    {
        "nome": "HS_Cotovelo_Velocidade_RMS_deg_s",
        "sub": "Graus/s (°/s)",
        "label": "Velocidade Angular RMS do Cotovelo",
        "unidade": "°/s",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Reatividade Articular",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 19.80
    },
    {
        "nome": "HS_Ombro_ROM_deg",
        "sub": "Graus (°)",
        "label": "Amplitude Total do Ângulo do Ombro (ROM)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Estratégia Motora",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 64.93
    },
    {
        "nome": "HS_Ombro_SD_deg",
        "sub": "Graus (°)",
        "label": "Variabilidade Postural do Ombro (Desvio Padrão)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Estratégia de Ombro",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 11.75
    },
    {
        "nome": "HS_Ombro_Velocidade_RMS_deg_s",
        "sub": "Graus/s (°/s)",
        "label": "Velocidade Angular RMS do Ombro",
        "unidade": "°/s",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Reatividade Articular",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 18.55
    },
    {
        "nome": "HS_Quadril_ROM_deg",
        "sub": "Graus (°)",
        "label": "Amplitude Total do Ângulo do Quadril (ROM)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Estratégia Motora",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 36.05
    },
    {
        "nome": "HS_Quadril_SD_deg",
        "sub": "Graus (°)",
        "label": "Variabilidade Postural do Quadril (Desvio Padrão)",
        "unidade": "Graus (°)",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Estratégia de Quadril",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 4.88
    },
    {
        "nome": "HS_Quadril_Velocidade_RMS_deg_s",
        "sub": "Graus/s (°/s)",
        "label": "Velocidade Angular RMS do Quadril",
        "unidade": "°/s",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Reatividade Articular",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 9.96
    },
    {
        "nome": "HS_Razao_Cotovelo_Ombro",
        "sub": "Razão SD (adimensional)",
        "label": "Razão de Variabilidade Postural (Cotovelo SD / Ombro SD)",
        "unidade": "Adimensional",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Assinatura Motora",
        "papel": "Fator Descritivo",
        "tipo_spss": "Escala",
        "p001_val": 1.44
    },
    {
        "nome": "HS_Estrategia_Dominante",
        "sub": "Cotovelo / Ombro / Punho",
        "label": "Classificação da Estratégia Motora Predominante",
        "unidade": "Texto",
        "bloco": "5. Handstand Estático (HS)",
        "cor_hex": "7030A0",
        "modelos": "Estatística Descritiva / Perfil de Equilíbrio",
        "papel": "Fator Descritivo",
        "tipo_spss": "Nominal",
        "p001_val": "Cotovelo-Dominante"
    },

    # --- BLOCO 6: BIOMECÂNICA HANDSTAND WALK (HSW - LaBioCoM) ---
    {
        "nome": "HSW_Distancia_Percorrida_Mediana_m",
        "sub": "Metros (m)",
        "label": "Mediana da Distância Percorrida no HSW (3 Tentativas)",
        "unidade": "m",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080", # Verde Petróleo / Teal
        "modelos": "DESFECHO PRINCIPAL HSW (Modelos 1 a 6 HSW)",
        "papel": "Variável Dependente (Desfecho Central)",
        "tipo_spss": "Escala",
        "p001_val": 1.78
    },
    {
        "nome": "HSW_Distancia_Percorrida_Max_m",
        "sub": "Metros (m)",
        "label": "Distância Máxima Percorrida em Tentativa Única",
        "unidade": "m",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "Estatística Descritiva HSW",
        "papel": "Descritiva",
        "tipo_spss": "Escala",
        "p001_val": 4.28
    },
    {
        "nome": "HSW_Variabilidade_Espacial_CV_pct",
        "sub": "CV (%)",
        "label": "Variabilidade Espacial do Comprimento de Passada (CV%)",
        "unidade": "%",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 2 (Eficiência da Locomoção)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 14.07
    },
    {
        "nome": "HSW_Variabilidade_Temporal_CV_pct",
        "sub": "CV (%)",
        "label": "Variabilidade Temporal do Tempo entre Passadas (CV%)",
        "unidade": "%",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 2 (Eficiência da Locomoção)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 56.95
    },
    {
        "nome": "HSW_Comprimento_Passada_Medio_mm",
        "sub": "Milímetros (mm)",
        "label": "Comprimento Médio de Passada no HSW",
        "unidade": "mm",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 2 (Eficiência da Locomoção)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 428.28
    },
    {
        "nome": "HSW_Cadencia_passos_s",
        "sub": "Passos / segundo",
        "label": "Cadência de Passadas no HSW",
        "unidade": "passos/s",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 2 (Eficiência da Locomoção)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 0.98
    },
    {
        "nome": "HSW_Indice_Verticalidade_deg",
        "sub": "Graus (°)",
        "label": "Índice de Verticalidade Dinâmico no HSW",
        "unidade": "Graus (°)",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 13.42
    },
    {
        "nome": "HSW_Extensao_Cervical_deg",
        "sub": "Graus (°)",
        "label": "Ângulo de Extensão Cervical Durante o Deslocamento",
        "unidade": "Graus (°)",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 42.64
    },
    {
        "nome": "HSW_Flexao_Plantar_deg",
        "sub": "Graus (°)",
        "label": "Ângulo de Flexão Plantar do Tornozelo no Deslocamento",
        "unidade": "Graus (°)",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 5.85
    },
    {
        "nome": "HSW_Base_Apoio_mm",
        "sub": "Milímetros (mm)",
        "label": "Largura Mediolateral Média da Base de Apoio no HSW",
        "unidade": "mm",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 566.77
    },
    {
        "nome": "HSW_Rotacao_Maos_deg",
        "sub": "Graus (°)",
        "label": "Orientação Angular das Mãos Durante a Locomoção",
        "unidade": "Graus (°)",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "HSW: Modelo 3 (Técnica)",
        "papel": "Preditor Independente",
        "tipo_spss": "Escala",
        "p001_val": 66.87
    },
    {
        "nome": "HSW_Velocidade_Media_m_s",
        "sub": "m/s",
        "label": "Velocidade Média de Deslocamento Linear",
        "unidade": "m/s",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "Estatística Descritiva / Comparação Grupos (Evita Colinearidade)",
        "papel": "Descritiva Especial",
        "tipo_spss": "Escala",
        "p001_val": 0.42
    },
    {
        "nome": "HSW_Tempo_Contato_Maos_s",
        "sub": "Segundos (s)",
        "label": "Tempo Médio de Contato Palmar por Apoio",
        "unidade": "s",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "Estatística Descritiva / Comparação Grupos (Evita Colinearidade)",
        "papel": "Descritiva Especial",
        "tipo_spss": "Escala",
        "p001_val": 1.05
    },
    {
        "nome": "HSW_Tempo_Duplo_Suporte_pct",
        "sub": "%",
        "label": "Percentual do Ciclo de Marcha em Duplo Suporte Palmar",
        "unidade": "%",
        "bloco": "6. Handstand Walk (HSW)",
        "cor_hex": "008080",
        "modelos": "Estatística Descritiva / Estratégia Motora",
        "papel": "Descritiva Especial",
        "tipo_spss": "Escala",
        "p001_val": 56.01
    },

    # --- BLOCO 7: METADADOS E ARQUIVOS ---
    {
        "nome": "Status_Coleta_LaCiDH",
        "sub": "Status",
        "label": "Status Operacional da Sessão 1 (LaCiDH)",
        "unidade": "Texto",
        "bloco": "7. Metadados e Arquivos",
        "cor_hex": "595959", # Cinza Escuro
        "modelos": "Gestão de Coleta",
        "papel": "Controle Operacional",
        "tipo_spss": "Nominal",
        "p001_val": "Concluída (LaCiDH)"
    },
    {
        "nome": "Status_Coleta_Biodex",
        "sub": "Status",
        "label": "Status da Dinamometria Isocinética Biodex",
        "unidade": "Texto",
        "bloco": "7. Metadados e Arquivos",
        "cor_hex": "595959",
        "modelos": "Gestão de Coleta",
        "papel": "Controle Operacional",
        "tipo_spss": "Nominal",
        "p001_val": "Concluída (Bilateral 70°)"
    },
    {
        "nome": "Status_Coleta_LaBioCoM",
        "sub": "Status",
        "label": "Status Operacional da Sessão 2 (LaBioCoM Vicon)",
        "unidade": "Texto",
        "bloco": "7. Metadados e Arquivos",
        "cor_hex": "595959",
        "modelos": "Gestão de Coleta",
        "papel": "Controle Operacional",
        "tipo_spss": "Nominal",
        "p001_val": "Concluída (Vicon + Bertec)"
    },
    {
        "nome": "Relatorio_BIA_PDF",
        "sub": "Nome do Arquivo",
        "label": "Arquivo PDF do Laudo Sanny BIA1011-AF",
        "unidade": "Texto",
        "bloco": "7. Metadados e Arquivos",
        "cor_hex": "595959",
        "modelos": "Arquivo Bruto",
        "papel": "Auditoria de Dados",
        "tipo_spss": "Nominal",
        "p001_val": "P001_sanny.pdf"
    },
    {
        "nome": "Relatorio_Biodex_PDF",
        "sub": "Nome do Arquivo",
        "label": "Arquivo PDF do Relatório Biodex Multi-Joint System",
        "unidade": "Texto",
        "bloco": "7. Metadados e Arquivos",
        "cor_hex": "595959",
        "modelos": "Arquivo Bruto",
        "papel": "Auditoria de Dados",
        "tipo_spss": "Nominal",
        "p001_val": "P001_biodex.pdf"
    },
    {
        "nome": "Observacoes_Gerais",
        "sub": "Texto",
        "label": "Observações Técnicas e Intercorrências",
        "unidade": "Texto",
        "bloco": "7. Metadados e Arquivos",
        "cor_hex": "595959",
        "modelos": "Controle Metodológico",
        "papel": "Auditoria de Dados",
        "tipo_spss": "Nominal",
        "p001_val": "Sessões 1 (LaCiDH) e 2 (LaBioCoM) concluídas com 100% de integridade. 9 trials C3D limpos e auditados sem falhas de marcadores."
    }
]


def gerar_tabelas_oficiais():
    base_raiz = r"C:\Users\Gui\Documents\MESTRADO_HANDSTAND\coleta\PILOTO OFICIAL"
    caminhos_saida = [
        os.path.join(base_raiz, "03_CONSOLIDACAO_DATASET", "dataset_final")
    ]

    for pasta in caminhos_saida:
        os.makedirs(pasta, exist_ok=True)

    print("=" * 85)
    print("GERAÇÃO DA TABELA OFICIAL DO DATASET DO MESTRADO (SPSS / ESTATÍSTICA FINAL)")
    print(f"Total de Variáveis Cadastradas: {len(VARIAVEIS_METRICAS)}")
    print("=" * 85)

    # 1. Gerar Arquivo Excel Formatado
    wb = openpyxl.Workbook()
    
    # --- ABA 1: DATASET PRINCIPAL (SPSS READY) ---
    ws_data = wb.active
    ws_data.title = "Dataset_Oficial_SPSS"
    ws_data.views.sheetView[0].showGridLines = True

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=9, italic=True, color="D9D9D9")
    font_data = Font(name="Calibri", size=11)
    font_pending = Font(name="Calibri", size=10, italic=True, color="7F7F7F")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    fill_pending = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    # Linha 1: Nome da Variável | Linha 2: Unidade/Subtítulo
    for col_idx, item in enumerate(VARIAVEIS_METRICAS, start=1):
        cor_header = item["cor_hex"]
        fill_head = PatternFill(start_color=cor_header, end_color=cor_header, fill_type="solid")
        
        c1 = ws_data.cell(row=1, column=col_idx, value=item["nome"])
        c1.fill = fill_head
        c1.font = font_header
        c1.alignment = align_center

        c2 = ws_data.cell(row=2, column=col_idx, value=item["sub"])
        c2.fill = fill_head
        c2.font = font_sub
        c2.alignment = align_center

    # Linha 3: Dados de P001
    for col_idx, item in enumerate(VARIAVEIS_METRICAS, start=1):
        val = item["p001_val"]
        c3 = ws_data.cell(row=3, column=col_idx, value=val)
        c3.border = border_thin
        
        if val == "":
            c3.fill = fill_pending
            c3.font = font_pending
        else:
            c3.font = font_data

        if isinstance(val, (int, float)):
            c3.alignment = align_right
        elif val in ["P001", "M", "F", "D", "E"]:
            c3.alignment = align_center
        else:
            c3.alignment = align_left

    # Auto-ajuste da largura das colunas
    for col in ws_data.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Congelar as primeiras duas linhas e a primeira coluna
    ws_data.freeze_panes = "B3"

    # --- ABA 2: DICIONÁRIO DE MODELOS E VARIÁVEIS ---
    ws_dict = wb.create_sheet(title="Dicionario_Modelos_Variaveis")
    ws_dict.views.sheetView[0].showGridLines = True

    headers_dict = [
        "Nome_Variavel_SPSS", "Rotulo_Descricao", "Unidade_Medida", "Dominio_Bloco",
        "Modelos_Estatisticos_Vinculados", "Papel_no_Modelo", "Tipo_Medida_SPSS"
    ]
    fill_dict_head = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

    for col_idx, h in enumerate(headers_dict, start=1):
        cell = ws_dict.cell(row=1, column=col_idx, value=h)
        cell.fill = fill_dict_head
        cell.font = font_header
        cell.alignment = align_center

    for row_idx, item in enumerate(VARIAVEIS_METRICAS, start=2):
        row_vals = [
            item["nome"], item["label"], item["unidade"], item["bloco"],
            item["modelos"], item["papel"], item["tipo_spss"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_dict.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_thin
            cell.alignment = align_left if col_idx != 1 else align_center

    for col in ws_dict.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_dict.column_dimensions[col_letter].width = max(max_len + 4, 16)

    ws_dict.freeze_panes = "B2"

    # Salvar nos diretórios de saída
    for pasta in caminhos_saida:
        caminho_xlsx = os.path.join(pasta, "TABELA_OFICIAL_MESTRADO_ESTATISTICA_FINAL.xlsx")
        caminho_csv = os.path.join(pasta, "TABELA_OFICIAL_MESTRADO_ESTATISTICA_FINAL.csv")

        wb.save(caminho_xlsx)
        print(f"[OK] Excel salvo em: {caminho_xlsx}")

        # Exportar CSV tabular com os dados (linha 1 = nomes, linha 2 = P001)
        with open(caminho_csv, "w", newline="", encoding="utf-8-sig") as f_csv:
            writer = csv.writer(f_csv)
            writer.writerow([item["nome"] for item in VARIAVEIS_METRICAS])
            writer.writerow([item["p001_val"] for item in VARIAVEIS_METRICAS])
        print(f"[OK] CSV salvo em:   {caminho_csv}")

    print("=" * 85)
    print("Processamento concluído com absoluto sucesso!")
    print("=" * 85)


if __name__ == "__main__":
    gerar_tabelas_oficiais()
