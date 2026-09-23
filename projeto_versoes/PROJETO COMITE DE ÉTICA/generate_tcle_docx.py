import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

doc = docx.Document()

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(12)

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p_title.add_run('TERMO DE CONSENTIMENTO LIVRE E ESCLARECIDO – TCLE\n(Conselho Nacional de Saúde, Resoluções 466/12 e 510/16)')
run.bold = True

doc.add_paragraph() # spacer

# Body paragraphs
p1 = doc.add_paragraph("O(a) Sr(a) está sendo convidado(a) a participar de uma pesquisa de Mestrado intitulada “ANÁLISE DE FATORES ASSOCIADOS AO DESEMPENHO NO HANDSTAND E HANDSTAND WALK”, que será desenvolvida por Guilherme de Paula Lemos, sob a orientação do Prof. Dr. Matheus Machado Gomes. O objetivo da referida pesquisa é investigar os fatores preditores do desempenho no handstand (parada de mãos) e no handstand walk (andar com as mãos) em praticantes de diferentes modalidades, integrando variáveis de força, controle postural, cinemática articular, antropometria, experiência prática e idade. Esta pesquisa visa responder como diferentes combinações dessas variáveis influenciam o desempenho na posição invertida, preenchendo uma lacuna na literatura científica e subsidiando programas de treinamento mais eficientes, seguros e específicos a cada modalidade. Esta pesquisa foi analisada e aprovada pelo Comitê de Ética em Pesquisa da Escola de Educação Física e Esporte de Ribeirão Preto, que tem como objetivo analisar as implicações éticas de projetos de pesquisa envolvendo seres humanos. Os endereços e contatos dos(as) pesquisadores(as) e do Comitê de Ética em Pesquisa estão disponíveis para o(a) Sr(a) no final deste documento.")
p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p2 = doc.add_paragraph("Antes de decidir se participará da pesquisa e assinar este termo, pedimos que leia as informações abaixo e esclareça quaisquer dúvidas que ainda tiver. É importante ressaltar que a sua participação nesta pesquisa será voluntária, e que o(a) Sr(a) poderá se recusar a participar ou retirar o seu consentimento a qualquer momento sem ser penalizado(a) por isso. Todas as informações obtidas na pesquisa serão mantidas em sigilo e sua identidade será preservada nos trabalhos científicos resultantes desta pesquisa. Ademais, ao final da pesquisa, você terá livre acesso aos seus dados individuais. Além disso, esclarecemos que eventuais despesas e danos decorrentes da sua participação nesta pesquisa serão ressarcidas e indenizados, respectivamente, quando necessário.")
p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p3 = doc.add_paragraph("Durante a sua participação na pesquisa, o(a) Sr(a) deverá comparecer a duas sessões de testes em dias distintos, com intervalo de até sete dias, nas dependências da Escola de Educação Física e Esporte de Ribeirão Preto (EEFERP – USP). Na primeira sessão (no Laboratório de Cineantropometria e Desempenho Humano), será realizada uma avaliação da composição corporal por meio de bioimpedância elétrica, o que exigirá um jejum prévio de 4 horas, abstenção de álcool e cafeína nas 24 horas antecedentes e esvaziamento da bexiga. Em seguida, o(a) Sr(a) realizará testes de resistência isométrica no handstand (com os pés escorados na parede), medição da força de preensão das mãos e dos dedos utilizando um dinamômetro, e um teste progressivo de carga máxima (1-RM) no exercício Shoulder Press (desenvolvimento de ombros). Na segunda sessão (no Laboratório de Biomecânica e Controle Motor), será analisado o seu desempenho no handstand livre e no handstand walk sobre uma plataforma de força, enquanto o seu movimento será filmado por um sistema tridimensional de captura de movimento, o que exigirá a fixação de marcadores esféricos reflexivos em diversos pontos do seu corpo. Ambas as sessões terão instruções prévias, aquecimento específico, e serão acompanhadas integralmente pelo pesquisador responsável. Cada sessão terá duração de aproximadamente 60 minutos.")
p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p4 = doc.add_paragraph("A participação na presente pesquisa não é livre de riscos. O procedimento envolvendo testes de esforço e de carga máxima pode gerar riscos inerentes à prática de exercícios físicos moderados e intensos, tais como desconfortos, fadiga, dor muscular tardia, além do risco de desequilíbrio, escorregões ou quedas durante as execuções do handstand. Para minimizar esses riscos, as avaliações serão conduzidas e supervisionadas presencialmente por um profissional de Educação Física capacitado e pelo pesquisador, os quais realizarão o monitoramento constante da intensidade e atuarão prontamente como apoio físico de segurança. Recomendamos que o(a) Sr(a) evite atividades físicas intensas nas 48 horas que antecedem cada avaliação. Caso os riscos venham a se concretizar (como uma entorse ou lesão muscular aguda), o pesquisador responsável prestará o suporte imediato de primeiros socorros e fará o encaminhamento ao atendimento médico na rede pública (pronto atendimento) ou particular de saúde, acompanhando o caso até a devida recuperação do participante. Embora apresente riscos, a participação na pesquisa também poderá gerar benefícios diretos ao participante, como conhecer de forma precisa os seus resultados de composição corporal, força muscular e parâmetros biomecânicos de controle e equilíbrio, que podem orientar sua prática esportiva. Além disso, a pesquisa trará o benefício indireto de contribuir para o avanço científico, auxiliando profissionais a desenvolverem métodos de ensino e treino mais eficazes para exercícios em inversão.")
p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p5 = doc.add_paragraph("Se, depois de solucionar suas dúvidas, o(a) Sr(a) se sentir esclarecido(a) sobre a pesquisa, seus objetivos, eventuais riscos e benefícios, o(a) convidamos a assinar duas vias deste termo, rubricando todas as suas páginas, sendo que uma via ficará com o(a) Sr(a) e a outra com o pesquisador responsável.")
p5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph("Ribeirão Preto, _____/_____/_____").alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph()

doc.add_paragraph("Dados sobre a Pesquisa:").runs[0].bold = True
p_proj = doc.add_paragraph("Título do Projeto: ")
p_proj.runs[0].bold = True
p_proj.add_run("ANÁLISE DE FATORES ASSOCIADOS AO DESEMPENHO NO HANDSTAND E HANDSTAND WALK")

doc.add_paragraph("Guilherme de Paula Lemos").runs[0].bold = True
doc.add_paragraph("Cargo/Função: Aluno de Mestrado\nInstituição: Escola de Educação Física e Esporte de Ribeirão Preto (EEFERP - USP)\nEndereço: Av. Bandeirantes, 3900 - Monte Alegre, 14040-907 - Ribeirão Preto SP\nTelefone: (16) 98122-2356\ne-mail: guilherme.lemos@usp.br")

doc.add_paragraph("Matheus Machado Gomes").runs[0].bold = True
doc.add_paragraph("Cargo/Função: Orientador Responsável / Professor Doutor\nInstituição: Escola de Educação Física e Esporte de Ribeirão Preto (EEFERP - USP)\nEndereço: Av. Bandeirantes, 3900 - Monte Alegre, 14040-907 - Ribeirão Preto SP\nTelefone: (16) 3315-8776\ne-mail: mmgomes@usp.br")

doc.add_paragraph()
doc.add_paragraph("Dados do(a) participante da pesquisa:").runs[0].bold = True
doc.add_paragraph("Nome completo:__________________________________________________________________________________")
doc.add_paragraph("Data de Nascimento: _____/_____/_____ Telefone para contato:___________________________________________")

doc.add_paragraph("\n")
doc.add_paragraph("________________________________________")
doc.add_paragraph("Assinatura do(a) Pesquisador(a) Responsável")

doc.add_paragraph("\n")
doc.add_paragraph("________________________________________")
doc.add_paragraph("Assinatura do(a) Participante")

doc.add_paragraph("\n\n")
p_comite = doc.add_paragraph()
r_comite = p_comite.add_run("Comitê de Ética em Pesquisa – Escola de Educação Física e Esporte de Ribeirão Preto\nAv. Bandeirantes, 3900 – Monte Alegre – CEP: 14040-907 – Ribeirão Preto – SP\nwww.eeferp.usp.br | cep90@usp.br | (16) 3315-0494\nAtendimento presencial: terças-feiras e quintas-feiras das 08:30 às 11:30.")
r_comite.font.size = Pt(10)

doc.save(r'TCLE\TCLE_Guilherme_Lemos_Atualizado.docx')
print("Document generated successfully.")
