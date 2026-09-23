#!/usr/bin/env python3
"""
Módulo unificado de comunicação (Gmail USP + WhatsApp Web) para a pesquisa de mestrado (Handstand - EEFERP/USP).
"""
import argparse
import openpyxl
import os
import subprocess
import urllib.parse
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
TABELA_MESTRE = BASE_DIR / "coleta" / "PILOTO OFICIAL" / "00_CADASTRO_TRIAGEM_E_TCLE" / "TABELA_MESTRE_IDENTIFICACAO_PARTICIPANTES.xlsx"
FIREFOX_PATH = r"C:\Program Files\Mozilla Firefox\firefox.exe"
DIR_ANEXOS = BASE_DIR / "automacoes" / "anexos_envio"

def carregar_participante(identificador):
    """Busca participante por ID (ex: P001) ou por nome na Tabela Mestre."""
    if not TABELA_MESTRE.exists():
        raise FileNotFoundError(f"Tabela Mestre não encontrada: {TABELA_MESTRE}")

    wb = openpyxl.load_workbook(TABELA_MESTRE, data_only=True)
    sheet = wb["02_Participantes_Oficiais"]
    
    headers = [cell for cell in next(sheet.iter_rows(values_only=True))]
    id_idx = headers.index("ID")
    nome_idx = headers.index("Nome_Completo")
    apelido_idx = headers.index("Nome_Abreviado")
    email_idx = headers.index("Email")
    tel_idx = headers.index("WhatsApp_Telefone")
    status_idx = headers.index("Status_Coleta")
    s1_idx = headers.index("Data_Sessao1_LaCiDH")
    s2_idx = headers.index("Data_Sessao2_LaBioCoM")

    identificador_clean = str(identificador).strip().lower()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        p_id = str(row[id_idx] or "").strip()
        p_nome = str(row[nome_idx] or "").strip()
        p_apelido = str(row[apelido_idx] or "").strip()
        
        if (p_id.lower() == identificador_clean or 
            identificador_clean in p_nome.lower() or 
            identificador_clean in p_apelido.lower()):
            
            tel = "".join(filter(str.isdigit, str(row[tel_idx] or "")))
            if len(tel) in [10, 11]:
                tel = f"55{tel}"

            return {
                "id": p_id,
                "nome_completo": p_nome,
                "nome_abreviado": p_apelido or p_nome,
                "email": row[email_idx],
                "whatsapp": tel,
                "status": row[status_idx],
                "sessao_1": row[s1_idx],
                "sessao_2": row[s2_idx]
            }

    return None

def abrir_no_firefox(url):
    if Path(FIREFOX_PATH).exists():
        cmd = f'Start-Process "{FIREFOX_PATH}" -ArgumentList "-new-tab", "`"{url}`""'
    else:
        cmd = f'Start-Process "`"{url}`""'
    subprocess.run(["powershell", "-NoProfile", "-Command", cmd], check=True)

def abrir_pasta_anexos():
    if DIR_ANEXOS.exists():
        cmd = f'Start-Process explorer.exe -ArgumentList "`"{DIR_ANEXOS}`""'
        subprocess.run(["powershell", "-NoProfile", "-Command", cmd], check=False)

def gerar_links_comunicacao(p_id, tipo_msg="devolutiva"):
    p = carregar_participante(p_id)
    if not p:
        raise ValueError(f"Participante {p_id} não encontrado na Tabela Mestre.")

    nome = p["nome_abreviado"]

    if tipo_msg == "devolutiva":
        assunto_email = f"[Pesquisa Mestrado Handstand USP] Envio do Relatório Individual Devolutivo - {nome}"
        corpo_email = f"""Olá {nome}, tudo bem?

Espero que este e-mail o encontre bem.

Gostaria de agradecer imensamente pela sua participação nas sessões de coleta da nossa pesquisa de Mestrado sobre o desempenho no Handstand e Handstand Walk, realizada na EEFERP-USP.

Conforme combinado, estou encaminhando o seu Relatório Individual Devolutivo em formato PDF, contendo os resultados detalhados das suas avaliações de força muscular isométrica e isocinética, taxas de desenvolvimento de torque, assimetrias bilaterais e controle postural.

Fique à vontade caso queira conversar sobre os dados ou tirar qualquer dúvida!

Atenciosamente,
Guilherme de Paula Lemos
Mestrando em Educação Física e Esporte — EEFERP/USP
Laboratório de Biomecânica e Controle Motor (LaBioCoM / LaCiDH)
guilherme.lemos@usp.br"""

        msg_wa = f"""Olá {nome}! Tudo bem? Aqui é o Guilherme da pesquisa de Mestrado da USP (Handstand).

Passando para agradecer novamente pela sua participação nas coletas da EEFERP-USP e avisar que o seu Relatório Individual Devolutivo já está pronto! 

Estou enviando o arquivo em anexo com os resultados das suas avaliações de força e estabilidade. Fique à vontade para tirar qualquer dúvida!"""

    else:
        assunto_email = f"[Pesquisa Mestrado Handstand USP] Contato - {nome}"
        corpo_email = f"Olá {nome},\n\nContato da pesquisa de Mestrado Handstand EEFERP-USP.\n\nAtenciosamente,\nGuilherme Lemos"
        msg_wa = f"Olá {nome}! Tudo bem? Aqui é o Guilherme da pesquisa de Mestrado da USP (Handstand)."

    # Link Gmail (/u/1/)
    params_gmail = {
        "view": "cm",
        "fs": "1",
        "to": p["email"],
        "su": assunto_email,
        "body": corpo_email
    }
    url_gmail = f"https://mail.google.com/mail/u/1/?{urllib.parse.urlencode(params_gmail, quote_via=urllib.parse.quote)}"

    # Link WhatsApp Web
    url_wa = f"https://web.whatsapp.com/send?phone={p['whatsapp']}&text={urllib.parse.quote(msg_wa)}"

    return {
        "participante": p,
        "email_link": url_gmail,
        "whatsapp_link": url_wa,
        "assunto_email": assunto_email,
        "corpo_email": corpo_email,
        "msg_whatsapp": msg_wa
    }

def main():
    parser = argparse.ArgumentParser(description="Disparador 1-clique para Gmail e WhatsApp")
    parser.add_argument("id", help="ID do participante (ex: P001, P002)")
    parser.add_argument("--tipo", default="devolutiva", choices=["devolutiva", "geral"])
    parser.add_argument("--abrir-email", action="store_true", help="Abre o Gmail no Firefox")
    parser.add_argument("--abrir-whatsapp", action="store_true", help="Abre o WhatsApp Web no Firefox")
    parser.add_argument("--abrir-anexos", action="store_true", help="Abre a pasta dos laudos em PDF")

    args = parser.parse_args()
    info = gerar_links_comunicacao(args.id, args.tipo)

    print(f"\nParticipante: {info['participante']['nome_completo']} ({info['participante']['id']})")
    print(f"E-mail: {info['participante']['email']}")
    print(f"WhatsApp: {info['participante']['whatsapp']}")
    print(f"\nLink Gmail: {info['email_link']}")
    print(f"Link WhatsApp: {info['whatsapp_link']}")

    if args.abrir_anexos:
        abrir_pasta_anexos()
    if args.abrir_email:
        abrir_no_firefox(info["email_link"])
    if args.abrir_whatsapp:
        abrir_no_firefox(info["whatsapp_link"])

if __name__ == "__main__":
    main()
